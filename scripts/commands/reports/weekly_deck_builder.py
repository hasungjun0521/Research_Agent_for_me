#!/usr/bin/env python3
"""PowerPoint assembly for the weekly dev deck.

Opens the user template (if any), fills placeholders, adds content slides from
the template's body layout, embeds rendered visuals, and strips decorative
empty-bracket shapes. Falls back to a built-in themed deck when no template.

python-pptx is imported at module load; callers must ensure it is installed.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from . import weekly_deck_charts as charts

SLIDE_W = Inches(10)
SLIDE_H = Inches(5.625)
_BRACKET_CHARS = {"[", "]", " ", "\n", " "}


# --- template helpers ------------------------------------------------------

def strip_bracket_decorations(prs) -> int:
    """Remove empty '[   ]' decorative shapes inherited from layouts."""
    removed = 0
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            for sh in list(layout.shapes):
                if sh.is_placeholder or not sh.has_text_frame:
                    continue
                text = sh.text_frame.text
                if text and set(text) <= _BRACKET_CHARS and ("[" in text or "]" in text):
                    sh._element.getparent().remove(sh._element)
                    removed += 1
    return removed


def _ph_by_idx(slide) -> dict:
    return {ph.placeholder_format.idx: ph for ph in slide.placeholders}


def _body_layout(prs):
    """Pick a 'title + body' layout for content slides (TITLE_ONLY_1-like)."""
    best = None
    for layout in prs.slide_layouts:
        idxs = {ph.placeholder_format.idx for ph in layout.placeholders}
        if {0, 1} <= idxs:
            name = (layout.name or "").upper()
            if "TITLE_ONLY" in name or best is None:
                best = layout
    return best or prs.slide_layouts[0]


def _fill_body(ph, items: list[str], header: str | None = None) -> None:
    tf = ph.text_frame
    tf.clear()
    first = tf.paragraphs[0]
    rows: list[tuple[str, bool, int | None]] = []
    if header is not None:
        rows.append((header, True, None))
    rows += [(it, False, 0) for it in items]
    if not rows:
        rows = [("이번 주 기록 없음", False, 0)]
    for i, (text, bold, level) in enumerate(rows):
        p = first if i == 0 else tf.add_paragraph()
        if level is not None:
            p.level = level
        run = p.add_run()
        run.text = text
        run.font.bold = bold
        if bold:
            run.font.size = Pt(15)


def _remove_body_placeholder(slide) -> None:
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx == 1:
            ph._element.getparent().remove(ph._element)


def _add_picture_fit(slide, image_path: Path, top_in: float = 0.95,
                     max_w_in: float = 8.4, max_h_in: float = 4.2) -> None:
    from PIL import Image  # Pillow ships with python-pptx's image stack
    try:
        with Image.open(image_path) as im:
            iw, ih = im.size
    except Exception:
        iw, ih = (max_w_in, max_h_in)
    ratio = iw / ih if ih else max_w_in / max_h_in
    w_in = max_w_in
    h_in = w_in / ratio
    if h_in > max_h_in:
        h_in = max_h_in
        w_in = h_in * ratio
    left = Inches((10 - w_in) / 2)
    slide.shapes.add_picture(str(image_path), left, Inches(top_in), width=Inches(w_in))


def _add_content_slide(prs, title: str):
    slide = prs.slides.add_slide(_body_layout(prs))
    _ph_by_idx(slide)[0].text = title
    return slide


# --- public build ----------------------------------------------------------

def build_deck(week: dict, template_path: Path | None, out_path: Path,
               build_dir: Path) -> Path:
    if template_path and Path(template_path).is_file():
        prs = _build_from_template(week, Path(template_path), build_dir)
    else:
        prs = _build_fallback(week, build_dir)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def _compose_visuals(week: dict, build_dir: Path) -> dict:
    """Render the shared visual assets once; return paths."""
    build_dir.mkdir(parents=True, exist_ok=True)
    stamp = week["until"].strftime("%Y%m%d")
    assets: dict = {}
    if week.get("kpis"):
        assets["kpis"] = charts.render_kpi_strip(week["kpis"], build_dir / f"kpis_{stamp}.png")
    if week.get("trend"):
        assets["trend"] = charts.render_trend_chart(
            week["trend"], build_dir / f"trend_{stamp}.png", ylabel=week.get("metric", "metric"))
    if week.get("deltas"):
        assets["deltas"] = charts.render_delta_bars(week["deltas"], build_dir / f"delta_{stamp}.png")
    return assets


def _build_from_template(week: dict, template_path: Path, build_dir: Path):
    prs = Presentation(str(template_path))
    strip_bracket_decorations(prs)
    assets = _compose_visuals(week, build_dir)
    slides = list(prs.slides)

    # cover (first slide) = project name + date
    cover = _ph_by_idx(slides[0])
    if 0 in cover:
        cover[0].text = week["title"]
    if 1 in cover:
        cover[1].text = week["cover_date"]
    for idx in (2, 3, 4):
        if idx in cover:
            cover[idx].text = ""

    # summary slide (reuse 2nd template slide if it is a body layout, else add)
    summary = slides[1] if len(slides) > 1 and {0, 1} <= set(_ph_by_idx(slides[1])) else _add_content_slide(prs, "")
    sp = _ph_by_idx(summary)
    sp[0].text = "이번 주 요약"
    _fill_body(sp[1], week.get("highlights", [])[:4], header=week.get("takeaway"))
    if assets.get("kpis"):
        summary.shapes.add_picture(str(assets["kpis"]), Inches(0.3), Inches(3.55), width=Inches(9.4))

    if assets.get("trend"):
        s = _add_content_slide(prs, f"실험 결과 추이 · {week.get('metric','metric')}")
        _remove_body_placeholder(s)
        _add_picture_fit(s, assets["trend"])
    if assets.get("deltas"):
        s = _add_content_slide(prs, "이번 주 변화 (Δ)")
        _remove_body_placeholder(s)
        _add_picture_fit(s, assets["deltas"])
    if week.get("figures"):
        _add_figures_slide_template(prs, week["figures"])

    plan = _add_content_slide(prs, "다음 주 액션")
    _fill_body(_ph_by_idx(plan)[1], week.get("next_actions", [])[:6])
    return prs


def _add_figures_slide_template(prs, figures: list[dict]) -> None:
    s = _add_content_slide(prs, "이번 주 산출물")
    _remove_body_placeholder(s)
    figs = figures[:4]
    cols = 2 if len(figs) > 1 else 1
    cell_w = 4.4 if cols == 2 else 8.0
    for i, fig in enumerate(figs):
        r, c = divmod(i, cols)
        left = Inches(0.6 + c * (cell_w + 0.4))
        top = Inches(1.05 + r * 2.0)
        try:
            s.shapes.add_picture(str(fig["path"]), left, top, width=Inches(cell_w))
        except Exception:
            continue


# --- fallback (no template) ------------------------------------------------

_NAVY = RGBColor(0x16, 0x1D, 0x2F)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_INK = RGBColor(0x1E, 0x29, 0x3B)
_ACC = RGBColor(0x4F, 0x46, 0xE5)


def _txt(slide, x, y, w, h, text, size, *, bold=False, color=_INK, align=None):
    from pptx.enum.text import PP_ALIGN
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align:
        p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _build_fallback(week: dict, build_dir: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    assets = _compose_visuals(week, build_dir)

    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, _NAVY)
    _rect(s, 0, 0, Inches(0.2), SLIDE_H, _ACC)
    _txt(s, Inches(0.7), Inches(2.0), Inches(8.8), Inches(1.1), week["title"], 34, bold=True, color=_WHITE)
    _txt(s, Inches(0.72), Inches(3.2), Inches(8.8), Inches(0.5), week["cover_date"], 16,
         color=RGBColor(0x94, 0xA3, 0xB8))

    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, SLIDE_W, Inches(0.75), _NAVY)
    _txt(s, Inches(0.4), Inches(0.12), Inches(9), Inches(0.5), "이번 주 요약", 22, bold=True, color=_WHITE)
    if week.get("takeaway"):
        _txt(s, Inches(0.5), Inches(0.95), Inches(9), Inches(0.5), week["takeaway"], 15, bold=True, color=_ACC)
    if assets.get("kpis"):
        s.shapes.add_picture(str(assets["kpis"]), Inches(0.3), Inches(1.5), width=Inches(9.4))
    y = 3.5
    for hl in week.get("highlights", [])[:4]:
        _txt(s, Inches(0.6), Inches(y), Inches(9), Inches(0.4), f"•  {hl}", 13)
        y += 0.42

    for key, title in (("trend", f"실험 결과 추이 · {week.get('metric','metric')}"), ("deltas", "이번 주 변화 (Δ)")):
        if assets.get(key):
            s = prs.slides.add_slide(blank)
            _rect(s, 0, 0, SLIDE_W, Inches(0.75), _NAVY)
            _txt(s, Inches(0.4), Inches(0.12), Inches(9), Inches(0.5), title, 22, bold=True, color=_WHITE)
            _add_picture_fit(s, assets[key])

    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, SLIDE_W, Inches(0.75), _NAVY)
    _txt(s, Inches(0.4), Inches(0.12), Inches(9), Inches(0.5), "다음 주 액션", 22, bold=True, color=_WHITE)
    y = 1.2
    actions = week.get("next_actions", [])[:6] or ["이번 주 기록 없음"]
    for a in actions:
        _txt(s, Inches(0.6), Inches(y), Inches(9), Inches(0.4), f"•  {a}", 14)
        y += 0.5
    return prs
